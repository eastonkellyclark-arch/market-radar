"""The deck's geometry, read back off the rendered file.

**Why this file exists separately from `test_decks.py`.** That file tests what a deck
*says* -- every substitution named, the cover leading with the worst input, no market
cap anywhere. This one tests where things physically are, which is a different failure
mode and one that no text assertion can see: python-pptx will place a box anywhere you
ask it to, including off the slide and on top of something else, and the file still
opens without complaint.

The pptx skill was not available in this environment, so these are the traps found by
rendering and reading back rather than the ones a skill would have listed. Three of
them were real and caught here first:

- normalising table widths to the page stretched the cover's half-width table under
  the data-quality panel
- the cover's two fixed-position footnotes collided as soon as a filer had six
  substitutions instead of two
- a regex that moved page bodies onto the grid also rewrote one textbox's *height*
  from 0.6 to the margin constant
"""

from __future__ import annotations

from datetime import date

import pytest

from marketradar import decks

pytest.importorskip("pptx", reason="python-pptx is in the `decks` group")

EPS = 0.02


def _subject(**over):
    base = dict(
        cik="0000000001", company="Fixture Industries Inc", ticker="FIX",
        sic=3011, first_period=date(2019, 3, 31), last_period=date(2026, 6, 30),
        still_filing=True,
        fundamentals={
            name: {"value": 1.0e9, "status": "stated", "tag": f"Tag{name}",
                   "coverage": 0.9}
            for name in ("revenue", "net_income", "assets", "liabilities",
                         "equity", "capex", "operating_cash_flow")
        },
        peers={"sic_depth": 2, "peers_banded": 24, "peers_material": 11,
               "turnover_median": 1.1, "turnover_iqr": 2.8},
        valuation={"enterprise_value": 1.2e9, "wacc": 0.11,
                   "terminal_share": 0.45, "beta": 1.3, "growth": 0.03,
                   "free_cash_flow": 9e7,
                   "flex": {0.0: 2.6e8, 0.03: 3.5e8, 0.06: 4.9e8, 0.10: 7.8e8},
                   # The heavy case, because it is the one that overflows: six
                   # caveats fill the cover panel and the footnote band.
                   "substitutions": ["peer_beta", "comp_depth_fallback",
                                     "absent_capex", "erp_constant",
                                     "growth_constant", "growth_mismatch"]},
        insiders=[], deals=[],
        # **Direction alternates on purpose.** The first version of this fixture rose
        # monotonically, so every candle closed up and the "both directions" test
        # failed on the fixture rather than on the code. A chart fixture that only
        # goes one way cannot tell a two-colour renderer from a one-colour one --
        # the same gap as a candle fixture that never spans a period boundary.
        prices=[(date(2024 + (m // 12), (m % 12) + 1, d), 10.0 + m, 11.5 + m,
                 8.5 + m, (10.0 + m) + (0.6 if m % 3 else -0.6), 1_000_000 + d)
                for m in range(30) for d in (3, 10, 17, 24)],
        unexplained_moves=4,
    )
    base.update(over)
    return decks.Subject(**base)


@pytest.fixture(scope="module")
def rendered(tmp_path_factory):
    from pptx import Presentation

    dest = decks.build(_subject(), tmp_path_factory.mktemp("d") / "deck.pptx")
    return Presentation(str(dest))


def _box(shape) -> tuple[float, float, float, float]:
    from pptx.util import Emu

    l, t = Emu(shape.left).inches, Emu(shape.top).inches
    return l, t, l + Emu(shape.width).inches, t + Emu(shape.height).inches


def _texts(slide):
    return [s for s in slide.shapes
            if s.has_text_frame and (s.text_frame.text or "").strip()]


def test_nothing_is_placed_off_the_slide(rendered) -> None:
    """python-pptx places a box wherever it is told, including past the edge, and
    saves a file that opens cleanly with the content simply not visible."""
    bad = []
    for n, slide in enumerate(rendered.slides, 1):
        for shape in slide.shapes:
            l, t, r, b = _box(shape)
            if l < -EPS or t < -EPS or r > decks.SLIDE_W + EPS \
                    or b > decks.SLIDE_H + EPS:
                bad.append(f"p{n}: ({l:.2f},{t:.2f})-({r:.2f},{b:.2f})")
    assert not bad, "shapes outside the slide:\n  " + "\n  ".join(bad)


def test_no_text_overlaps_other_text(rendered) -> None:
    """Two of the three layout defects in this rebuild were overlaps, and neither
    was visible in the text of the file -- both decks read correctly and looked
    wrong. Rules and panels are excluded: they are meant to sit under things."""
    bad = []
    for n, slide in enumerate(rendered.slides, 1):
        boxes = _texts(slide)
        for i, a in enumerate(boxes):
            al, at, ar, ab = _box(a)
            for other in boxes[i + 1:]:
                bl, bt, br, bb = _box(other)
                ox = min(ar, br) - max(al, bl)
                oy = min(ab, bb) - max(at, bt)
                if ox > 0.25 and oy > 0.12:
                    bad.append(
                        f"p{n}: {ox:.2f}x{oy:.2f}in between "
                        f"{a.text_frame.text[:24]!r} and "
                        f"{other.text_frame.text[:24]!r}")
    assert not bad, "overlapping text:\n  " + "\n  ".join(bad)


def test_every_page_has_the_footnote_band_and_content_stays_above_it(
        rendered) -> None:
    """**The constraint that does not move, as geometry.**

    `test_decks.py` pins that the band's *words* are on every page. This pins that
    the band has a place of its own and that no page's content is sitting on top of
    it -- which is the difference between a designed footnote band and a caveat
    textbox dropped wherever there was room.
    """
    for n, slide in enumerate(rendered.slides, 1):
        band = [s for s in slide.shapes
                if abs(_box(s)[1] - decks.BAND_TOP) < EPS
                and _box(s)[2] - _box(s)[0] > decks.SLIDE_W - 0.1]
        assert band, f"p{n} has no footnote band"
        for shape in _texts(slide):
            _l, t, _r, b = _box(shape)
            assert not (t < decks.BAND_TOP - EPS and b > decks.BAND_TOP + 0.12), (
                f"p{n}: content crosses into the band: "
                f"{shape.text_frame.text[:40]!r}")


def test_every_size_is_on_the_type_scale(rendered) -> None:
    """A scale with exceptions is not a scale. Eight pages were still asking for
    11, 13 and 14pt after the primitives were converted, which is exactly the
    "size=13 here, size=12 there" the named constants were introduced to stop."""
    scale = {decks.MICRO, decks.SMALL, decks.BODY, decks.H2, decks.LEAD,
             decks.H1, decks.COVER_NAME, decks.DISPLAY}
    seen: set[float] = set()
    for slide in rendered.slides:
        for shape in _texts(slide):
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size is not None:
                        seen.add(run.font.size.pt)
    off = sorted(s for s in seen if s not in scale)
    assert not off, f"sizes not on the scale: {off} (scale is {sorted(scale)})"


def test_only_the_two_declared_faces_are_used(rendered) -> None:
    """A third font is either an accident or a theme leaking through, and a deck is
    the artifact that leaves the room -- it gets opened on a machine that resolves
    whatever it is told to."""
    faces = set()
    for slide in rendered.slides:
        for shape in _texts(slide):
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.name:
                        faces.add(run.font.name)
    assert faces <= {decks.FONT_DISPLAY, decks.FONT_BODY}, faces
    assert faces, "no font was set anywhere, so the template's theme decides"


def test_the_chart_draws_candles_in_both_directions(rendered) -> None:
    """**Shapes, and the dashboard's own two hues.** A chart page that rendered as
    a table of summary numbers -- which is what this page used to be -- would pass
    every text assertion in the other file."""
    prices = rendered.slides[len(decks.PAGES) - 1]
    hues: dict[str, int] = {}
    for shape in prices.shapes:
        try:
            if shape.fill.type is None:
                continue
            rgb = str(shape.fill.fore_color.rgb)
        except Exception:
            continue
        hues[rgb] = hues.get(rgb, 0) + 1
    assert hues.get(decks.CANDLE_UP, 0) > 0, "no up candles"
    assert hues.get(decks.CANDLE_DOWN, 0) > 0, "no down candles"
    # Body, wick and volume bar per candle, so the count is a multiple of the bars
    # rather than one mark each.
    assert hues[decks.CANDLE_UP] + hues[decks.CANDLE_DOWN] >= 30, hues


def test_a_subject_with_no_prices_says_so_rather_than_drawing_nothing(
        tmp_path) -> None:
    """The heavy archetype in the real population is a delisted filer with no
    resolvable ticker, so this is the common case rather than an edge one."""
    from pptx import Presentation

    dest = decks.build(_subject(prices=[]), tmp_path / "d.pptx")
    page = Presentation(str(dest)).slides[len(decks.PAGES) - 1]
    text = "\n".join(s.text_frame.text for s in _texts(page))
    assert "No price history" in text
    assert "survivor-only" in text
